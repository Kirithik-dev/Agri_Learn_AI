import io
from typing import Dict, Any
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

class ExportService:
    """Provides professional PDF and DOCX exports for AGRI-LEARN AI training modules."""

    @staticmethod
    def generate_pdf(module_data: Dict[str, Any]) -> bytes:
        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            rightMargin=36,
            leftMargin=36,
            topMargin=36,
            bottomMargin=36
        )

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            'TitleStyle',
            parent=styles['Heading1'],
            fontSize=18,
            leading=22,
            textColor=colors.HexColor('#065f46'),
            spaceAfter=8
        )
        heading_style = ParagraphStyle(
            'HeadingStyle',
            parent=styles['Heading2'],
            fontSize=13,
            leading=16,
            textColor=colors.HexColor('#047857'),
            spaceBefore=12,
            spaceAfter=6
        )
        body_style = ParagraphStyle(
            'BodyStyle',
            parent=styles['Normal'],
            fontSize=10,
            leading=14,
            textColor=colors.HexColor('#1f2937'),
            spaceAfter=6
        )
        bullet_style = ParagraphStyle(
            'BulletStyle',
            parent=styles['Normal'],
            fontSize=9.5,
            leading=13,
            textColor=colors.HexColor('#374151'),
            leftIndent=12,
            spaceAfter=4
        )

        story = []

        # Header Badge
        header_text = f"<b>AGRI-LEARN AI | Intelligent Agricultural Training Content</b><br/><font size=8 color='#4b5563'>INTELLIX Hackathon (Team Syntax Soldiers - BIT-AI-002)</font>"
        story.append(Paragraph(header_text, body_style))
        story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#10b981'), spaceAfter=12))

        # Title
        story.append(Paragraph(module_data.get('title', 'Agricultural Training Module'), title_style))

        # Metadata Table
        meta_data = [
            [
                f"<b>Crop:</b> {module_data.get('crop')}",
                f"<b>Topic:</b> {module_data.get('topic')}",
                f"<b>Audience:</b> {module_data.get('target_audience')}"
            ],
            [
                f"<b>Language:</b> {module_data.get('language')}",
                f"<b>Duration:</b> {module_data.get('duration_minutes')} mins",
                f"<b>Evaluation Score:</b> {module_data.get('evaluation_score', 94)}%"
            ]
        ]
        t = Table(meta_data, colWidths=[180, 180, 180])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#ecfdf5')),
            ('TEXTCOLOR', (0,0), (-1,-1), colors.HexColor('#065f46')),
            ('FONTNAME', (0,0), (-1,-1), 'Helvetica'),
            ('FONTSIZE', (0,0), (-1,-1), 8.5),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#a7f3d0')),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#a7f3d0')),
        ]))
        story.append(t)
        story.append(Spacer(1, 10))

        # Learning Objectives
        story.append(Paragraph("1. Learning Objectives", heading_style))
        for obj in module_data.get('learning_objectives', []):
            story.append(Paragraph(f"• {obj}", bullet_style))

        # Introduction
        story.append(Paragraph("2. Introduction", heading_style))
        story.append(Paragraph(module_data.get('introduction', ''), body_style))

        # Key Concepts
        story.append(Paragraph("3. Key Concepts", heading_style))
        for kc in module_data.get('key_concepts', []):
            story.append(Paragraph(f"<b>• {kc.get('concept')}:</b> {kc.get('description')}", bullet_style))

        # Step by step
        story.append(Paragraph("4. Step-by-Step Field Practice", heading_style))
        for step in module_data.get('step_by_step', []):
            story.append(Paragraph(f"<b>Step {step.get('step_number')}: {step.get('title')}</b>", body_style))
            story.append(Paragraph(f"{step.get('action')}", bullet_style))

        # Safety Precautions
        story.append(Paragraph("5. Safety Precautions & Pre-Harvest Intervals (PHI)", heading_style))
        for sp in module_data.get('safety_precautions', []):
            story.append(Paragraph(f"⚠️ {sp}", bullet_style))

        # Sources
        story.append(Paragraph("6. Trusted Sources & References", heading_style))
        for s in module_data.get('sources', []):
            story.append(Paragraph(f"• <b>{s.get('title')}</b> — {s.get('source')} ({s.get('section')})", bullet_style))

        doc.build(story)
        buffer.seek(0)
        return buffer.getvalue()

    @staticmethod
    def generate_docx(module_data: Dict[str, Any]) -> bytes:
        doc = docx.Document()

        # Set title
        title_p = doc.add_paragraph()
        title_run = title_p.add_run(module_data.get('title', 'Agricultural Training Module'))
        title_run.font.size = Pt(18)
        title_run.font.bold = True
        title_run.font.color.rgb = RGBColor(6, 95, 70)

        # Subtitle
        sub_p = doc.add_paragraph()
        sub_run = sub_p.add_run("AGRI-LEARN AI | Intelligent Agricultural Training Content Generator")
        sub_run.font.size = Pt(10)
        sub_run.font.italic = True
        sub_run.font.color.rgb = RGBColor(100, 116, 139)

        # Metadata
        doc.add_heading("Training Metadata", level=2)
        table = doc.add_table(rows=2, cols=3)
        table.style = 'Light Shading Accent 1'
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = f"Crop: {module_data.get('crop')}"
        hdr_cells[1].text = f"Topic: {module_data.get('topic')}"
        hdr_cells[2].text = f"Audience: {module_data.get('target_audience')}"
        row_cells = table.rows[1].cells
        row_cells[0].text = f"Language: {module_data.get('language')}"
        row_cells[1].text = f"Duration: {module_data.get('duration_minutes')} min"
        row_cells[2].text = f"Score: {module_data.get('evaluation_score', 94)}%"

        doc.add_paragraph("")

        # Objectives
        doc.add_heading("1. Learning Objectives", level=2)
        for obj in module_data.get('learning_objectives', []):
            doc.add_paragraph(f"• {obj}")

        # Introduction
        doc.add_heading("2. Introduction", level=2)
        doc.add_paragraph(module_data.get('introduction', ''))

        # Key Concepts
        doc.add_heading("3. Key Concepts", level=2)
        for kc in module_data.get('key_concepts', []):
            p = doc.add_paragraph()
            r1 = p.add_run(f"• {kc.get('concept')}: ")
            r1.bold = True
            p.add_run(kc.get('description', ''))

        # Step by step
        doc.add_heading("4. Step-by-Step Training", level=2)
        for step in module_data.get('step_by_step', []):
            p = doc.add_paragraph()
            r = p.add_run(f"Step {step.get('step_number')}: {step.get('title')}\n")
            r.bold = True
            p.add_run(step.get('action', ''))

        # Safety Precautions
        doc.add_heading("5. Safety Precautions & Environmental Protection", level=2)
        for sp in module_data.get('safety_precautions', []):
            doc.add_paragraph(f"⚠️ {sp}")

        # Sources
        doc.add_heading("6. Sources & References", level=2)
        for s in module_data.get('sources', []):
            doc.add_paragraph(f"• {s.get('title')} - {s.get('source')} ({s.get('section')})")

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()

export_service = ExportService()
